from __future__ import annotations

import base64
import json
import os
import tempfile
import unittest
import zipfile
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread
from unittest.mock import patch

from evals.gdpval.load_instances import (
    load_instances,
    read_task_ids_file,
)
from evals.gdpval.judge_suite import GdpvalGsbJudgeSuite, GdpvalJudgeSuite
from evals.gdpval.suite import GdpvalSuite
from simple_agent_lab.evals import (
    RESULT_KEY,
    TRACE_KEY,
    LocalDirStore,
    LocalProcessBackend,
    RunArtifacts,
    run_suite_instance,
)
from simple_agent_lab.evals.suites.gdpval import (
    container,
    judge_excel_tools,
    judge_container,
    judge_gsb_container,
    judge_mcp,
)
from simple_agent_lab.evals.suites.gdpval.judge_scoring import (
    normalize_rubrics,
    parse_gsb_direction_payload,
    parse_gsb_judge_payload,
    parse_judge_payload,
    score_gsb_judgment,
    score_judgment,
)
from simple_agent_lab.evals.suites.gdpval.judge_mcp import (
    gdpval_mcp_server_configs,
    is_gdpval_judge_read_only_mcp_tool_name,
    normalize_judge_tool_mode,
)
from simple_agent_lab.evals.suites.gdpval.prompts import (
    GDPVAL_SYSTEM_PROMPT,
    gdpval_system_prompt,
)
from simple_agent_lab.evals.suites.gdpval.tools import make_gdpval_tools
from simple_agent_lab import Agent, assistant_message
from simple_agent_lab.compression import maybe_compress_context
from simple_agent_lab.llm.provider import Provider
from simple_agent_lab.messages import (
    ImageBlock,
    TextBlock,
    ToolCallBlock,
    ToolResultBlock,
    tool_results_message,
    tool_results_of,
)
from simple_agent_lab.state import State
from simple_agent_lab.tools import AgentTool, text_result, tool_result_text


class _QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, format: str, *args) -> None:
        return None


class GdpvalSuiteTest(unittest.TestCase):
    def test_task_input_hides_gold_and_inlines_reference_blobs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "refs").mkdir()
            (root / "refs" / "brief.txt").write_text("reference text", encoding="utf-8")

            suite = GdpvalSuite(reference_root=root)
            visible = suite.task_input(
                {
                    "task_id": "task-1",
                    "prompt": "Create a deliverable.",
                    "reference_files": ["refs/brief.txt"],
                    "deliverable_files": ["gold-answer.txt"],
                    "rubric_json": [{"score": 1, "criterion": "secret"}],
                }
            )

            self.assertEqual(visible["solver_agent_mode"], "tool-call-context-managed")
            self.assertTrue(visible["enable_web_tools"])
            self.assertNotIn("deliverable_files", visible)
            self.assertNotIn("rubric_json", visible)
            self.assertEqual(visible["reference_files"][0]["name"], "refs/brief.txt")
            self.assertNotIn("data_base64", visible["reference_files"][0])
            decoded = base64.b64decode(
                visible["reference_file_blobs"][0]["data_base64"]
            ).decode("utf-8")
            self.assertEqual(decoded, "reference text")

    def test_task_input_threads_solver_web_tool_flag(self) -> None:
        visible = GdpvalSuite(enable_web_tools=False).task_input(
            {
                "task_id": "task-1",
                "prompt": "Create a deliverable.",
            }
        )

        self.assertFalse(visible["enable_web_tools"])

    def test_task_input_uses_official_reference_files_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            task_dir = root / "task-1"
            task_dir.mkdir()
            (task_dir / "Population v2.xlsx").write_bytes(b"xlsx bytes")

            visible = GdpvalSuite(reference_root=root).task_input(
                {
                    "task_id": "task-1",
                    "prompt": "Create a deliverable.",
                    "reference_file_urls": [
                        "https://example.invalid/reference_files/hash/Population%20v2.xlsx"
                    ],
                    "reference_files": ["reference_files/hash/Population v2.xlsx"],
                }
            )

            blob = visible["reference_file_blobs"][0]
            self.assertFalse(blob.get("missing"))
            self.assertEqual(base64.b64decode(blob["data_base64"]), b"xlsx bytes")

    def test_task_input_downloads_reference_urls_without_reference_root(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "Population v2.xlsx").write_bytes(b"downloaded bytes")
            server = ThreadingHTTPServer(
                ("127.0.0.1", 0),
                lambda *args, **kwargs: _QuietHandler(
                    *args, directory=str(root), **kwargs
                ),
            )
            thread = Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                port = server.server_address[1]
                visible = GdpvalSuite().task_input(
                    {
                        "task_id": "task-1",
                        "prompt": "Create a deliverable.",
                        "reference_files": ["reference_files/hash/Population v2.xlsx"],
                        "reference_file_urls": [
                            f"http://127.0.0.1:{port}/Population%20v2.xlsx"
                        ],
                    }
                )
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)

            blob = visible["reference_file_blobs"][0]
            self.assertFalse(blob.get("missing"))
            self.assertEqual(base64.b64decode(blob["data_base64"]), b"downloaded bytes")

    def test_load_instances_filters_empty_deliverable_files_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rows.jsonl"
            path.write_text(
                "\n".join(
                    [
                        json.dumps(
                            {
                                "task_id": "empty",
                                "prompt": "p",
                                "deliverable_files": [],
                            }
                        ),
                        json.dumps(
                            {
                                "task_id": "nonempty",
                                "prompt": "p",
                                "deliverable_files": ["deliverable_files/x/out.xlsx"],
                            }
                        ),
                    ]
                )
                + "\n",
                encoding="utf-8",
            )

            rows = load_instances(path)

            self.assertEqual([row["task_id"] for row in rows], ["nonempty"])

    def test_load_instances_can_include_empty_deliverable_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rows.jsonl"
            path.write_text(
                "\n".join(
                    [
                        json.dumps(
                            {
                                "task_id": "empty",
                                "prompt": "p",
                                "deliverable_files": [],
                            }
                        ),
                        json.dumps(
                            {
                                "task_id": "nonempty",
                                "prompt": "p",
                                "deliverable_files": ["deliverable_files/x/out.xlsx"],
                            }
                        ),
                    ]
                )
                + "\n",
                encoding="utf-8",
            )

            rows = load_instances(path, require_deliverables=False)

            self.assertEqual([row["task_id"] for row in rows], ["empty", "nonempty"])

    def test_read_task_ids_file_skips_comments_and_deduplicates(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "task_ids.txt"
            path.write_text(
                "\n".join(
                    [
                        "# comment",
                        "task-a",
                        "task-b, task-c",
                        "task-a",
                        "",
                    ]
                ),
                encoding="utf-8",
            )

            self.assertEqual(
                read_task_ids_file(path),
                ["task-a", "task-b", "task-c"],
            )

    def test_prepare_writes_reference_files_and_task_mentions_paths(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workdir = Path(tmp) / "task" / "workdir"
            blob = base64.b64encode(b"reference text").decode("ascii")
            instance = {
                "task_id": "task-1",
                "prompt": "Use the reference.",
                "reference_files": [{"name": "brief.txt", "source": "brief.txt"}],
                "reference_file_blobs": [
                    {"name": "brief.txt", "source": "brief.txt", "data_base64": blob}
                ],
            }

            context = container.prepare(workdir, instance)
            reference_dir = Path(context["reference_dir"])
            self.assertEqual(
                (reference_dir / "brief.txt").read_text(encoding="utf-8"),
                "reference text",
            )

            task = container.build_task(instance, workdir=str(workdir))
            self.assertIn(f"WORKDIR: {workdir}", task)
            self.assertIn(str(reference_dir / "brief.txt"), task)
            self.assertIn("<FINAL_ANSWER>", GDPVAL_SYSTEM_PROMPT)
            self.assertIn("</FINAL_ANSWER>", GDPVAL_SYSTEM_PROMPT)
            self.assertIn("execute_bash as the primary tool", GDPVAL_SYSTEM_PROMPT)
            self.assertIn("multi_edit_file", GDPVAL_SYSTEM_PROMPT)
            self.assertIn("view_image", GDPVAL_SYSTEM_PROMPT)
            self.assertIn("Use WebSearch only", GDPVAL_SYSTEM_PROMPT)
            self.assertIn("Use WebFetch for a known public", GDPVAL_SYSTEM_PROMPT)
            self.assertNotIn("ImageSearch", GDPVAL_SYSTEM_PROMPT)
            self.assertIn(
                "External web tools are not available",
                gdpval_system_prompt(enable_web_tools=False),
            )
            self.assertIn("<FINAL_ANSWER>...</FINAL_ANSWER>", task)

    def test_gdpval_solver_tools_cover_bash_edit_todo_and_images(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workdir = Path(tmp) / "workdir"
            reference_dir = Path(tmp) / "reference_task_id_files"
            tools = {
                tool.name: tool
                for tool in make_gdpval_tools(
                    workdir=workdir,
                    reference_dir=reference_dir,
                    enable_web_tools=False,
                    output_head_chars=2000,
                    output_tail_chars=2000,
                )
            }
            self.assertEqual(
                set(tools),
                {"execute_bash", "TodoWrite", "multi_edit_file", "view_image"},
            )

            def call(name: str, args: dict) -> str:
                result = tools[name].execute("call-1", args, lambda: False, None)
                self.assertFalse(result.is_error, tool_result_text(result))
                return tool_result_text(result)

            call(
                "execute_bash",
                {"command": "printf 'alpha\\nbeta\\n' > answer.txt"},
            )
            self.assertIn(
                '"status": "pending"',
                call(
                    "TodoWrite",
                    {
                        "todos": [
                            {
                                "content": "create answer",
                                "status": "pending",
                            }
                        ]
                    },
                ),
            )
            self.assertIn(
                "gamma",
                call(
                    "multi_edit_file",
                    {
                        "file_path": str(workdir / "answer.txt"),
                        "edits": [
                            {
                                "old_string": "beta\n",
                                "new_string": "gamma\n",
                            }
                        ],
                    },
                ),
            )
            self.assertEqual(
                (workdir / "answer.txt").read_text(encoding="utf-8"),
                "alpha\ngamma\n",
            )

            failed = tools["multi_edit_file"].execute(
                "call-2",
                {
                    "file_path": str(workdir / "answer.txt"),
                    "edits": [
                        {
                            "old_string": "missing",
                            "new_string": "should-not-write",
                        }
                    ],
                },
                lambda: False,
                None,
            )
            self.assertTrue(failed.is_error)
            self.assertEqual(
                (workdir / "answer.txt").read_text(encoding="utf-8"),
                "alpha\ngamma\n",
            )

            pixel_png = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
                "/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
            )
            (workdir / "pixel.png").write_bytes(pixel_png)
            self.assertIn("Image file:", call("view_image", {"path": "pixel.png"}))

    def test_gdpval_solver_tools_include_jina_web_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workdir = Path(tmp) / "workdir"
            reference_dir = Path(tmp) / "reference_task_id_files"
            tools = {
                tool.name: tool
                for tool in make_gdpval_tools(
                    workdir=workdir,
                    reference_dir=reference_dir,
                )
            }

            self.assertEqual(
                set(tools),
                {
                    "execute_bash",
                    "TodoWrite",
                    "multi_edit_file",
                    "view_image",
                    "WebSearch",
                    "WebFetch",
                },
            )

    def test_gdpval_web_search_uses_jina_and_filters_domains(self) -> None:
        class _Response:
            status = 200
            headers = {"Content-Type": "application/json"}

            def __enter__(self):
                return self

            def __exit__(self, *exc: object) -> None:
                return None

            def read(self) -> bytes:
                return json.dumps(
                    {
                        "code": 200,
                        "status": 200,
                        "data": [
                            {
                                "title": "Keep",
                                "url": "https://example.com/a",
                                "description": "alpha",
                                "content": "# Alpha\n\nSearch result content.",
                            },
                            {
                                "title": "Drop",
                                "url": "https://blocked.test/a",
                                "description": "beta",
                                "content": "Blocked content.",
                            },
                        ],
                    }
                ).encode("utf-8")

        with tempfile.TemporaryDirectory() as tmp:
            tools = {
                tool.name: tool
                for tool in make_gdpval_tools(
                    workdir=Path(tmp) / "workdir",
                    reference_dir=Path(tmp) / "reference_task_id_files",
                )
            }
            with patch.dict(os.environ, {"JINA_API_KEY": "secret"}, clear=True):
                with patch(
                    "simple_agent_lab.evals.suites.gdpval.web_tools.urllib.request.urlopen",
                    return_value=_Response(),
                ) as urlopen:
                    result = tools["WebSearch"].execute(
                        "call-1",
                        {
                            "query": "alpha",
                            "max_results": 5,
                            "allowed_domains": ["example.com"],
                            "blocked_domains": ["blocked.test"],
                        },
                        lambda: False,
                        None,
                    )

            self.assertFalse(result.is_error, tool_result_text(result))
            payload = json.loads(tool_result_text(result))
            self.assertTrue(payload["ok"])
            self.assertEqual(payload["provider"], "jina")
            self.assertEqual(
                [item["link"] for item in payload["organic"]],
                ["https://example.com/a"],
            )
            self.assertEqual(urlopen.call_args.kwargs["timeout"], 60)
            self.assertIn("https://s.jina.ai/", urlopen.call_args.args[0].full_url)
            self.assertIn("site%3Aexample.com", urlopen.call_args.args[0].full_url)
            self.assertIn(
                "Search result content",
                payload["organic"][0]["content_preview"],
            )

    def test_gdpval_web_search_retries_timeout(self) -> None:
        class _Response:
            status = 200
            headers = {"Content-Type": "application/json"}

            def __enter__(self):
                return self

            def __exit__(self, *exc: object) -> None:
                return None

            def read(self) -> bytes:
                return json.dumps(
                    {
                        "code": 200,
                        "status": 200,
                        "data": [
                            {
                                "title": "Recovered",
                                "url": "https://example.com/recovered",
                                "description": "retry worked",
                                "content": "Recovered after timeout.",
                            }
                        ],
                    }
                ).encode("utf-8")

        with tempfile.TemporaryDirectory() as tmp:
            tools = {
                tool.name: tool
                for tool in make_gdpval_tools(
                    workdir=Path(tmp) / "workdir",
                    reference_dir=Path(tmp) / "reference_task_id_files",
                )
            }
            with patch.dict(os.environ, {"JINA_API_KEY": "secret"}, clear=True):
                with patch(
                    "simple_agent_lab.evals.suites.gdpval.web_tools.time.sleep"
                ) as sleep:
                    with patch(
                        "simple_agent_lab.evals.suites.gdpval.web_tools.urllib.request.urlopen",
                        side_effect=[TimeoutError("slow"), _Response()],
                    ) as urlopen:
                        result = tools["WebSearch"].execute(
                            "call-1",
                            {"query": "alpha", "max_results": 5},
                            lambda: False,
                            None,
                        )

            self.assertFalse(result.is_error, tool_result_text(result))
            payload = json.loads(tool_result_text(result))
            self.assertTrue(payload["ok"])
            self.assertEqual(urlopen.call_count, 2)
            self.assertEqual(urlopen.call_args.kwargs["timeout"], 60)
            sleep.assert_called_once_with(1)
            self.assertEqual(
                payload["organic"][0]["link"],
                "https://example.com/recovered",
            )

    def test_gdpval_web_fetch_uses_jina_and_writes_cache(self) -> None:
        class _Response:
            status = 200
            headers = {"Content-Type": "text/plain"}

            def __enter__(self):
                return self

            def __exit__(self, *exc: object) -> None:
                return None

            def read(self) -> bytes:
                return (
                    b"Title: Example Page\n"
                    b"URL Source: https://example.com/page\n"
                    b"Markdown Content:\n"
                    b"# Heading\n\nBody text from Jina."
                )

        with tempfile.TemporaryDirectory() as tmp:
            workdir = Path(tmp) / "workdir"
            tools = {
                tool.name: tool
                for tool in make_gdpval_tools(
                    workdir=workdir,
                    reference_dir=Path(tmp) / "reference_task_id_files",
                )
            }
            with patch.dict(os.environ, {"JINA_API_KEY": "jina-key"}, clear=True):
                with patch(
                    "simple_agent_lab.evals.suites.gdpval.web_tools.urllib.request.urlopen",
                    return_value=_Response(),
                ) as urlopen:
                    result = tools["WebFetch"].execute(
                        "call-1",
                        {"url": "https://example.com/page", "max_chars": 1000},
                        lambda: False,
                        None,
                    )

            self.assertFalse(result.is_error, tool_result_text(result))
            payload = json.loads(tool_result_text(result))
            self.assertTrue(payload["ok"])
            self.assertEqual(payload["provider"], "jina")
            self.assertEqual(payload["title"], "Example Page")
            self.assertIn("# Heading", payload["content"])
            content_path = Path(payload["content_path"])
            metadata_path = Path(payload["metadata_path"])
            self.assertTrue(content_path.is_file())
            self.assertTrue(metadata_path.is_file())
            self.assertIn(".webfetch_cache", content_path.parts)
            self.assertEqual(
                urlopen.call_args.args[0].full_url,
                "https://r.jina.ai/https://example.com/page",
            )

    def test_extract_result_collects_manifest_and_archive(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workdir = Path(tmp) / "workdir"
            workdir.mkdir()
            (workdir / "answer.txt").write_text("answer", encoding="utf-8")
            cache = workdir / ".webfetch_cache" / "abc"
            cache.mkdir(parents=True)
            (cache / "content.md").write_text("cached web text", encoding="utf-8")

            result = container.extract_result(workdir, {"task_id": "task-1"})

            self.assertEqual(result["status"], "solver_finished")
            self.assertEqual(result["files"][0]["relative_path"], "answer.txt")
            self.assertEqual(len(result["files"]), 1)
            self.assertGreater(result["workspace_archive_bytes"], 0)
            self.assertTrue(Path(result["workspace_archive_path"]).is_file())

    def test_gdpval_judges_install_image_once_context_policy(self) -> None:
        provider = Provider(id="fake", api="fake", model="fake-model")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            solver = container.build_agent(provider=provider, cwd=root / "solver")
            judge_context = {
                "input_dir": str(root / "judge" / "judge_inputs"),
            }
            gsb_context = {
                "input_dir": str(root / "gsb_judge" / "judge_inputs"),
            }
            judge_instance = {"judge_tool_mode": "local"}
            with judge_container.agent_context(
                provider=provider,
                cwd=root / "judge",
                instance=judge_instance,
                context=judge_context,
            ) as judge:
                rubric_tool_names = [tool.name for tool in judge.tools]
                self.assertIsNotNone(judge.context_policy)
                self.assertIsNotNone(judge.context_policy.strategy)
            with judge_gsb_container.agent_context(
                provider=provider,
                cwd=root / "gsb_judge",
                instance=judge_instance,
                context=gsb_context,
            ) as gsb_judge:
                gsb_tool_names = [tool.name for tool in gsb_judge.tools]
                self.assertIsNotNone(gsb_judge.context_policy)
                self.assertIsNotNone(gsb_judge.context_policy.strategy)

        self.assertIsNone(solver.context_policy)
        self.assertEqual(
            [tool.name for tool in solver.tools],
            [
                "execute_bash",
                "TodoWrite",
                "multi_edit_file",
                "view_image",
                "WebSearch",
                "WebFetch",
            ],
        )
        self.assertNotIn("read_file", {tool.name for tool in solver.tools})
        self.assertNotIn("write_file", {tool.name for tool in solver.tools})
        self.assertEqual(rubric_tool_names, gsb_tool_names)
        self.assertNotIn("write_file", rubric_tool_names)
        self.assertNotIn("execute_bash", rubric_tool_names)
        self.assertNotIn("read_file", rubric_tool_names)
        self.assertIn("excel_profile_sheet", rubric_tool_names)

    def test_gdpval_judge_image_context_policy_keeps_image_for_one_turn(
        self,
    ) -> None:
        def idle(visible):
            del visible
            return assistant_message("done", sender="judge", target="user")

        state = State("judge")
        state.send("task", "user", "judge", state.task)
        state.record(
            assistant_message(
                [TextBlock("read pdf"), ToolCallBlock("c0", "pdf_read_pdf", {})],
                sender="judge",
                target="user",
                kind="step",
            )
        )
        state.record(
            tool_results_message(
                [
                    ToolResultBlock(
                        tool_call_id="c0",
                        tool_name="pdf_read_pdf",
                        content=(
                            TextBlock("page text"),
                            ImageBlock(data="QUJD", mime_type="image/png"),
                        ),
                    )
                ],
                target="judge",
            )
        )

        policy = judge_mcp.gdpval_judge_context_policy()
        assert policy.strategy is not None
        self.assertIsNone(policy.strategy(state.active_context_items(), "judge"))

        state.record(
            assistant_message(
                "I inspected the image.",
                sender="judge",
                target="user",
                kind="step",
            )
        )
        events = maybe_compress_context(Agent("judge", idle), state, policy)
        self.assertEqual(events[1].strategy, "gdpval-judge-image-once")

        active_result_messages = [
            message
            for message in state.active_context_messages()
            if message.kind == "tool_result"
        ]
        self.assertEqual(len(active_result_messages), 1)
        active_results = tool_results_of(active_result_messages[0].content)
        self.assertFalse(
            any(
                isinstance(block, ImageBlock)
                for result in active_results
                for block in result.content
            )
        )
        self.assertIn(
            "Image omitted after one GDPVal judge model turn",
            active_results[0].content[1].text,
        )

        original_results = tool_results_of(state.messages[2].content)
        self.assertTrue(
            any(
                isinstance(block, ImageBlock)
                for result in original_results
                for block in result.content
            )
        )

        state.record(
            assistant_message(
                [TextBlock("read again"), ToolCallBlock("c1", "pdf_read_pdf", {})],
                sender="judge",
                target="user",
                kind="step",
            )
        )
        state.record(
            tool_results_message(
                [
                    ToolResultBlock(
                        tool_call_id="c1",
                        tool_name="pdf_read_pdf",
                        content=(ImageBlock(data="REVG", mime_type="image/png"),),
                    )
                ],
                target="judge",
            )
        )
        self.assertIsNone(policy.strategy(state.active_context_items(), "judge"))

    def test_judge_suite_threads_tool_mode_into_instance(self) -> None:
        suite = GdpvalGsbJudgeSuite(judge_tool_mode="local")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run_dir = root / "candidate"
            run_dir.mkdir()
            artifacts = RunArtifacts(
                instance_id="task-1",
                run_dir=run_dir,
                trajectory_path=run_dir / "out" / "trajectory.jsonl",
                status_code=0,
            )
            instance = suite.build_instance(
                {"task_id": "task-1", "prompt": "p"},
                candidate_result={"status": "solver_finished", "files": []},
                candidate_artifacts=artifacts,
            )

        self.assertEqual(instance["judge_tool_mode"], "local")

    def test_judge_mcp_registry_uses_local_stdio_commands(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            try:
                configs = gdpval_mcp_server_configs(
                    workdir=root / "workdir",
                    reference_dir=root / "judge_inputs",
                )
            except ModuleNotFoundError:
                self.skipTest("mcp extra not installed")

        by_name = {config.name: config for config in configs}
        self.assertEqual(by_name["filesystem"].command, "mcp-server-filesystem")
        self.assertEqual(by_name["pdf"].command, "pdf-reader-mcp")
        self.assertEqual(by_name["excel"].command, "excel-mcp-server")
        self.assertEqual(by_name["excel"].args, ("stdio",))
        self.assertEqual(by_name["word"].command, "word_mcp_server")
        self.assertEqual(by_name["ppt"].command, "ppt_mcp_server")

    def test_judge_mcp_filter_matches_swalm_read_only_keyword_rules(self) -> None:
        allowed = [
            "read_file",
            "list_directory",
            "get_workbook_metadata",
            "extract_slide_text",
            "search_files",
            "inspect_pdf",
            "parse_document",
            "profile_sheet",
            "query_table",
            "analyze_workbook",
            "describe_content",
            "metadata_info",
        ]
        denied = [
            "write_file",
            "edit_file",
            "create_directory",
            "move_file",
            "write_data_to_excel",
            "create_workbook",
            "delete_worksheet",
            "create_document",
            "add_paragraph",
            "search_and_replace",
            "save_document",
            "export_pdf",
            "import_document",
            "render_page",
        ]

        for name in allowed:
            with self.subTest(name=name):
                self.assertTrue(is_gdpval_judge_read_only_mcp_tool_name(name))
        for name in denied:
            with self.subTest(name=name):
                self.assertFalse(is_gdpval_judge_read_only_mcp_tool_name(name))

        # The swalm rule is intentionally keyword-based, not an exact-name list.
        self.assertFalse(
            is_gdpval_judge_read_only_mcp_tool_name("validate_formula_syntax")
        )
        self.assertFalse(is_gdpval_judge_read_only_mcp_tool_name("get_merged_cells"))

    def test_judge_mcp_filter_allows_ppt_stateful_read_helpers(self) -> None:
        allowed = [
            "open_presentation",
            "load_presentation",
            "import_presentation",
            "render_slide",
            "export_slide_image",
            "get_thumbnail",
            "extract_slide_text",
        ]
        denied = [
            "save_presentation",
            "edit_slide",
            "format_text",
            "add_slide",
            "replace_image",
        ]

        for name in allowed:
            with self.subTest(name=name):
                self.assertTrue(
                    is_gdpval_judge_read_only_mcp_tool_name(
                        name,
                        server_name="ppt",
                    )
                )
                self.assertTrue(
                    is_gdpval_judge_read_only_mcp_tool_name(
                        name,
                        server_name="ppt_mcp_server",
                    )
                )
        for name in denied:
            with self.subTest(name=name):
                self.assertFalse(
                    is_gdpval_judge_read_only_mcp_tool_name(
                        name,
                        server_name="ppt",
                    )
                )

    def test_judge_tool_mode_normalization(self) -> None:
        self.assertEqual(normalize_judge_tool_mode(None), "hybrid")
        self.assertEqual(normalize_judge_tool_mode("MCP"), "mcp")
        with self.assertRaises(ValueError):
            normalize_judge_tool_mode("remote")

    def test_gsb_judge_local_tool_surface_excludes_solver_bash(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with judge_mcp.open_gdpval_judge_tools(
                workdir=root / "workdir",
                reference_dir=root / "judge_inputs",
                mode="local",
                include_excel_helpers=True,
            ) as tools:
                tool_names = {tool.name for tool in tools}

        self.assertIn("excel_profile_sheet", tool_names)
        self.assertIn("read_data_from_excel_compact", tool_names)
        self.assertNotIn("execute_bash", tool_names)
        self.assertNotIn("read_file", tool_names)
        self.assertNotIn("write_file", tool_names)

    def test_gsb_judge_prompt_prefers_read_only_filesystem_tools(self) -> None:
        prompt = judge_gsb_container.GDPVAL_GSB_JUDGE_SYSTEM_PROMPT

        self.assertIn("senior industry expert", prompt)
        self.assertIn("File deliverables take precedence", prompt)
        self.assertIn("Special note", prompt)
        self.assertIn("filesystem_* read-only tools", prompt)
        self.assertIn("never output", prompt)
        self.assertIn('a bare "error" string', prompt)
        self.assertNotIn('simply output "error"', prompt)
        self.assertIn("Do not wrap the JSON in Markdown code fences", prompt)
        self.assertNotIn("bash", prompt.lower())
        self.assertNotIn("python libraries", prompt.lower())

    def test_judge_mcp_repairs_ab_label_paths_by_direction(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workspace = root / "workdir"
            reference = root / "judge_inputs"
            gold = reference / "gold"
            candidate = reference / "candidate"
            gold.mkdir(parents=True)
            candidate.mkdir(parents=True)
            workspace.mkdir()
            gold_file = gold / "report.txt"
            candidate_file = candidate / "report.txt"
            gold_file.write_text("gold", encoding="utf-8")
            candidate_file.write_text("candidate", encoding="utf-8")
            seen_args: dict[str, str] = {}

            def execute(call_id, args, abort, on_update):
                del call_id, abort, on_update
                seen_args.update(args)
                return text_result("ok")

            tool = judge_mcp._sanitize_judge_tool_output(
                AgentTool(
                    name="filesystem_read_file",
                    description="read",
                    parameters={"type": "object"},
                    execute=execute,
                ),
                workspace=workspace,
                reference_dir=reference,
                path_label_roles={"A": "gold", "B": "candidate"},
            )

            result = tool.execute(
                "call-1", {"path": "A/report.txt"}, lambda: False, None
            )

        self.assertFalse(result.is_error)
        self.assertEqual(Path(seen_args["path"]), gold_file.resolve())

    def test_judge_mcp_word_fallback_extracts_docx_text(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workspace = root / "workdir"
            reference = root / "judge_inputs"
            workspace.mkdir()
            reference.mkdir()
            docx_path = reference / "answer.docx"
            with zipfile.ZipFile(docx_path, "w") as archive:
                archive.writestr(
                    "word/document.xml",
                    (
                        "<w:document><w:body><w:p><w:r><w:t>"
                        "Fallback paragraph"
                        "</w:t></w:r></w:p></w:body></w:document>"
                    ),
                )

            def execute(call_id, args, abort, on_update):
                del call_id, args, abort, on_update
                return text_result("attributes construct error", is_error=True)

            tool = judge_mcp._sanitize_judge_tool_output(
                AgentTool(
                    name="word_get_document_text",
                    description="word text",
                    parameters={"type": "object"},
                    execute=execute,
                ),
                workspace=workspace,
                reference_dir=reference,
            )

            result = tool.execute(
                "call-1",
                {"document_path": str(docx_path)},
                lambda: False,
                None,
            )

        self.assertFalse(result.is_error)
        text = tool_result_text(result)
        self.assertIn("Fallback DOCX extraction succeeded", text)
        self.assertIn("Fallback paragraph", text)

    def test_judge_mcp_compacts_large_notebook_text(self) -> None:
        notebook = {
            "nbformat": 4,
            "cells": [
                {
                    "cell_type": "code",
                    "source": ["print('hello')\n", "x = '" + ("a" * 170_000) + "'"],
                    "outputs": [
                        {
                            "output_type": "stream",
                            "text": "hello\n" + ("b" * 10_000),
                        }
                    ],
                }
            ],
        }

        processed = judge_mcp._preprocess_judge_tool_text(json.dumps(notebook))

        self.assertIn("Notebook compacted", processed)
        self.assertLess(len(processed), 130_000)

    def test_judge_excel_tools_auto_correct_header_and_token_match_sheet(
        self,
    ) -> None:
        try:
            from openpyxl import Workbook
        except ModuleNotFoundError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "book.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Annual Revenue Report"
            sheet["A1"] = "generated report title"
            sheet["A2"] = "Entity"
            sheet["B2"] = "Amount"
            sheet["A3"] = "Acme"
            sheet["B3"] = 10
            workbook.save(workbook_path)

            filtered = judge_excel_tools._run_excel_action(
                "filter_rows",
                {
                    "filepath": str(workbook_path),
                    "sheet_name": "Revenue Report Annual",
                    "header_row": 1,
                    "filters": {"Entity": "Acme"},
                    "columns": ["Amount"],
                },
                workspace=root,
                references=root,
            )
            aggregated = judge_excel_tools._run_excel_action(
                "aggregate",
                {
                    "filepath": str(workbook_path),
                    "sheet_name": "Revenue Report Annual",
                    "header_row": 1,
                    "metrics": [
                        {"column": "Amount", "op": "sum", "name": "sum_amount"}
                    ],
                },
                workspace=root,
                references=root,
            )

        self.assertEqual(filtered["sheet_name"], "Annual Revenue Report")
        self.assertEqual(filtered["requested_header_row"], 1)
        self.assertEqual(filtered["header_row"], 2)
        self.assertEqual(filtered["header_auto_correction"]["to_header_row"], 2)
        self.assertEqual(filtered["rows"][0]["Amount"], 10)
        self.assertEqual(aggregated["header_row"], 2)
        self.assertEqual(aggregated["groups"][0]["sum_amount"], 10.0)

    def test_judge_excel_path_resolution_matches_swalm_ab_priority(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workdir = root / "workdir"
            deliverable = root / "deliverable_task_id_files"
            reference = root / "reference_task_id_files"
            workdir.mkdir()
            deliverable.mkdir()
            reference.mkdir()
            (workdir / "report.xlsx").write_bytes(b"candidate")
            (deliverable / "report.xlsx").write_bytes(b"standard")
            (reference / "report.xlsx").write_bytes(b"reference")

            a_path = judge_excel_tools._resolve_workbook_path(
                {"filepath": "A/report.xlsx"},
                workspace=root,
                references=root,
            )
            b_path = judge_excel_tools._resolve_workbook_path(
                {"filepath": "B/report.xlsx"},
                workspace=root,
                references=root,
            )

        self.assertEqual(a_path, deliverable / "report.xlsx")
        self.assertEqual(b_path, workdir / "report.xlsx")

    def test_judge_mcp_compacts_large_excel_cells_payload(self) -> None:
        raw = json.dumps(
            {
                "filepath": "/workspace/report.xlsx",
                "sheet_name": "Sheet1",
                "cells": [
                    {"address": "A1", "value": "Name"},
                    {"address": "B1", "value": "Amount"},
                    {"address": "A2", "value": "Acme"},
                    {"address": "B2", "value": 123.0},
                ],
            }
        )

        compacted = judge_mcp._maybe_compact_large_excel_cells_payload(
            raw,
            min_chars=0,
        )

        self.assertIsNotNone(compacted)
        assert compacted is not None
        self.assertIn("reason=large_excel_cells_output", compacted)
        self.assertIn("header_or_first_row", compacted)
        self.assertIn('row 2: ["Acme", "123.0"]', compacted)

    def test_judge_scoring_normalizes_rubrics_and_weighted_score(self) -> None:
        rubrics = json.dumps(
            [
                {"score": 2, "criterion": "Has workbook"},
                {"weight": 3, "rubric_content": "Reconciles totals"},
            ]
        )
        payload = parse_judge_payload(
            """
            ```json
            {
              "rubric_results": [
                {"index": 0, "grade": 1, "explanation": "present"},
                {"index": 1, "grade": 0.5, "explanation": "partial"}
              ],
              "overall_explanation": "mixed"
            }
            ```
            """
        )

        result = score_judgment(payload, rubrics)

        self.assertEqual(
            [item["criterion"] for item in normalize_rubrics(rubrics)],
            ["Has workbook", "Reconciles totals"],
        )
        self.assertEqual(result["status"], "judged")
        self.assertAlmostEqual(result["earned_score"], 3.5)
        self.assertAlmostEqual(result["max_score"], 5.0)
        self.assertAlmostEqual(result["score"], 0.7)

    def test_judge_parser_accepts_judge_result_tags(self) -> None:
        payload = parse_judge_payload(
            """
            <judge_result>
            {
              "rubric_results": [
                {"index": 0, "grade": 1, "explanation": "ok"}
              ],
              "overall_explanation": "done"
            }
            </judge_result>
            """
        )

        self.assertEqual(payload["rubric_results"][0]["grade"], 1)
        self.assertEqual(payload["overall_explanation"], "done")

    def test_gsb_judge_scoring_uses_forward_and_reverse_maps(self) -> None:
        rubrics = [{"score": 1, "criterion": "Candidate is better"}]
        payload = parse_gsb_judge_payload(
            """
            ```json
            {
              "reverse": {
                "rubrics_result": [
                  {
                    "index": 0,
                    "grade_A": 0,
                    "grade_B": 1,
                    "gsb": "A<B",
                    "grade_explanation": "candidate wins"
                  }
                ],
                "overall": {
                  "overall_explanation": "candidate wins",
                  "final_gsb": "A<B"
                }
              },
              "forward": {
                "rubrics_result": [
                  {
                    "index": 0,
                    "grade_A": 1,
                    "grade_B": 0,
                    "gsb": "A>B",
                    "grade_explanation": "candidate wins"
                  }
                ],
                "overall": {
                  "overall_explanation": "candidate wins",
                  "final_gsb": "A>B"
                }
              }
            }
            ```
            """
        )

        result = score_gsb_judgment(payload, rubrics)

        self.assertEqual(result["status"], "gsb_judged")
        self.assertAlmostEqual(result["rubrics_weighted_score_reverse"], 1.0)
        self.assertAlmostEqual(result["rubrics_weighted_score_forward"], 1.0)
        self.assertAlmostEqual(result["combined_weighted_score"], 1.0)
        self.assertAlmostEqual(result["llm_score"], 1.0)
        self.assertAlmostEqual(result["score_process"], 1.0)
        self.assertAlmostEqual(result["dcg_winrate"], 1.0)
        self.assertAlmostEqual(result["score"], 1.0)

    def test_gsb_judge_scoring_falls_back_to_raw_overall_when_no_rubrics(
        self,
    ) -> None:
        result = score_gsb_judgment(
            {
                "reverse": {
                    "rubrics_result": [],
                    "overall": {
                        "overall_explanation": "candidate is better",
                        "final_gsb": "A<B",
                    },
                }
            },
            [],
        )

        self.assertEqual(result["status"], "no_rubrics")
        self.assertEqual(result["score"], 0.75)

    def test_gsb_direction_parser_accepts_swalm_xmlish_output(self) -> None:
        payload = parse_gsb_direction_payload(
            """
            <rubrics_result>
              <rubric score="2" criterion="Accuracy" grade_A="1"
                grade_B="0" gsb="A>B">
                <grade_explanation>standard is more accurate</grade_explanation>
              </rubric>
            </rubrics_result>
            <overall>
              <overall_explanation>standard wins</overall_explanation>
              <final_gsb>A>B</final_gsb>
            </overall>
            """
        )

        self.assertEqual(payload["rubrics_result"][0]["score"], 2.0)
        self.assertEqual(payload["rubrics_result"][0]["criterion"], "Accuracy")
        self.assertEqual(payload["rubrics_result"][0]["gsb"], "A>B")
        self.assertEqual(payload["overall"]["final_gsb"], "A>B")

    def test_gsb_last_assistant_text_keeps_full_response_for_parsing(self) -> None:
        response = (
            "<rubrics_result>"
            + json.dumps(
                [
                    {
                        "score": 1,
                        "criterion": "Accuracy",
                        "grade_A": 1,
                        "grade_B": 0,
                        "gsb": "A>B",
                        "grade_explanation": "standard is better",
                    }
                ]
            )
            + "</rubrics_result>"
            + (" filler" * 40)
            + '<overall>{"overall_explanation":"standard wins",'
            + '"final_gsb":"A>B"}</overall>'
        )
        self.assertGreater(len(response), 120)
        state = State(task="")
        state.send(
            "final",
            sender="gdpval_gsb_judge_reverse_1",
            target="user",
            content=response,
        )

        text = judge_gsb_container._last_assistant_text(state)
        payload = parse_gsb_direction_payload(text)

        self.assertEqual(text, response)
        self.assertEqual(payload["overall"]["final_gsb"], "A>B")
        self.assertEqual(len(payload["rubrics_result"]), 1)

    def test_gsb_direction_attempt_limit_defaults_to_one_and_is_bounded(self) -> None:
        env = judge_gsb_container.GSB_DIRECTION_ATTEMPTS_ENV
        with patch.dict(os.environ, {}, clear=False):
            os.environ.pop(env, None)
            self.assertEqual(judge_gsb_container._max_gsb_direction_attempts({}), 1)
            self.assertEqual(
                judge_gsb_container._max_gsb_direction_attempts(
                    {"judge_gsb_direction_attempts": 3}
                ),
                3,
            )
            self.assertEqual(
                judge_gsb_container._max_gsb_direction_attempts(
                    {"judge_gsb_direction_attempts": 99}
                ),
                judge_gsb_container.MAX_GSB_DIRECTION_ATTEMPTS,
            )
            self.assertEqual(
                judge_gsb_container._max_gsb_direction_attempts(
                    {"judge_gsb_direction_attempts": 0}
                ),
                1,
            )

        with patch.dict(os.environ, {env: "2"}):
            self.assertEqual(judge_gsb_container._max_gsb_direction_attempts({}), 2)

    def test_gsb_attempt_result_fields_include_attempt_details(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workdir = Path(tmp)
            attempts = [
                {
                    "direction": "reverse",
                    "attempt": 1,
                    "has_tool_messages": True,
                    "mcp_tool_count": 2,
                    "failure_reason": "invalid_output",
                    "raw_response_preview": "preview",
                }
            ]
            (workdir / judge_gsb_container.JUDGE_ATTEMPTS_FILE).write_text(
                json.dumps({"attempts": attempts}),
                encoding="utf-8",
            )

            fields = judge_gsb_container._attempt_result_fields(workdir)

        self.assertEqual(fields["judge_attempts"], attempts)
        self.assertEqual(fields["judge_retry_summary"]["reverse"]["attempts"], 1)
        self.assertEqual(
            fields["judge_retry_summary"]["reverse"]["last_failure_reason"],
            "invalid_output",
        )

    def test_gsb_empty_direction_output_is_invalid_not_an_exception(self) -> None:
        payload, reason = judge_gsb_container._parse_direction_payload_safely("")

        self.assertEqual(payload, {})
        self.assertIn("invalid_output", reason)
        self.assertIn("GSB direction payload is empty", reason)

    def test_gsb_payload_attempt_failure_marks_result_invalid(self) -> None:
        reason = judge_gsb_container._payload_failure_reason(
            {
                "_sal_judge_attempt_summary": {
                    "reverse": {
                        "attempts": 1,
                        "last_failure_reason": (
                            "invalid_output: ValueError: GSB direction payload is empty"
                        ),
                    }
                }
            }
        )

        self.assertIn("reverse judge failed", reason)
        self.assertIn("GSB direction payload is empty", reason)

    def test_gsb_scoring_counts_missing_direction_like_swalm(self) -> None:
        result = score_gsb_judgment(
            {
                "reverse": {
                    "rubrics_result": [
                        {"index": 0, "grade_A": 0, "grade_B": 1, "gsb": "A<B"},
                        {"index": 1, "grade_A": 1, "grade_B": 0, "gsb": "A>B"},
                    ],
                    "overall": {"final_gsb": "A<B"},
                }
            },
            [
                {"score": 103, "criterion": "large candidate win"},
                {"score": 97, "criterion": "small candidate loss"},
            ],
        )

        self.assertAlmostEqual(result["rubrics_weighted_score_reverse"], 0.03)
        self.assertAlmostEqual(result["combined_weighted_score_raw"], 0.015)
        self.assertEqual(result["combined_weighted_score"], 0.5)
        self.assertEqual(result["score"], 0.5)

    def test_judge_prepare_stages_candidate_gold_and_references(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate_workdir = root / "candidate_workdir"
            candidate_workdir.mkdir()
            (candidate_workdir / "answer.txt").write_text("answer", encoding="utf-8")
            candidate_result = container.extract_result(
                candidate_workdir, {"task_id": "task-1"}
            )
            deliverable_root = root / "deliverables"
            reference_root = root / "references"
            (deliverable_root / "gold").mkdir(parents=True)
            (reference_root / "refs").mkdir(parents=True)
            (deliverable_root / "gold" / "answer.txt").write_text(
                "gold", encoding="utf-8"
            )
            (reference_root / "refs" / "brief.txt").write_text(
                "brief", encoding="utf-8"
            )
            suite = GdpvalJudgeSuite(
                image="unused",
                deliverable_root=deliverable_root,
                reference_root=reference_root,
                network_mode=None,
            )
            judge_instance = suite.build_instance(
                {
                    "task_id": "task-1",
                    "prompt": "Make answer.",
                    "rubrics": [{"score": 1, "criterion": "Answer exists"}],
                    "deliverable_files": ["gold/answer.txt"],
                    "reference_files": ["refs/brief.txt"],
                },
                candidate_result=candidate_result,
                candidate_artifacts=RunArtifacts(
                    instance_id="task-1",
                    run_dir=root / "solver-run",
                    trajectory_path=root / "solver-run" / "out" / "trajectory.jsonl",
                    status_code=0,
                ),
            )

            judge_workdir = root / "judge" / "workdir"
            context = judge_container.prepare(judge_workdir, judge_instance)
            task = judge_container.build_task(
                judge_instance, workdir=str(judge_workdir)
            )

            candidate_dir = Path(context["candidate_dir"])
            gold_dir = Path(context["gold_dir"])
            reference_dir = Path(context["reference_dir"])
            self.assertEqual(
                (candidate_dir / "workdir" / "answer.txt").read_text(encoding="utf-8"),
                "answer",
            )
            self.assertEqual(
                (gold_dir / "gold" / "answer.txt").read_text(encoding="utf-8"),
                "gold",
            )
            self.assertEqual(
                (reference_dir / "refs" / "brief.txt").read_text(encoding="utf-8"),
                "brief",
            )
            self.assertNotIn(judge_container.JUDGE_RESULT_FILE, task)
            self.assertIn("<judge_result>", task)
            self.assertIn("Do not write files", task)

    def test_gsb_judge_prepare_and_task_use_gold_as_standard_answer(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate_workdir = root / "candidate_workdir"
            candidate_workdir.mkdir()
            (candidate_workdir / "answer.txt").write_text("answer", encoding="utf-8")
            candidate_result = container.extract_result(
                candidate_workdir, {"task_id": "task-1"}
            )
            deliverable_root = root / "deliverables"
            (deliverable_root / "gold").mkdir(parents=True)
            (deliverable_root / "gold" / "answer.txt").write_text(
                "gold", encoding="utf-8"
            )
            suite = GdpvalGsbJudgeSuite(
                image="unused",
                deliverable_root=deliverable_root,
                network_mode=None,
            )
            judge_instance = suite.build_instance(
                {
                    "task_id": "task-1",
                    "prompt": "Make answer.",
                    "rubrics": [{"score": 1, "criterion": "Answer quality"}],
                    "deliverable_files": ["gold/answer.txt"],
                },
                candidate_result=candidate_result,
                candidate_artifacts=RunArtifacts(
                    instance_id="task-1",
                    run_dir=root / "solver-run",
                    trajectory_path=root / "solver-run" / "out" / "trajectory.jsonl",
                    status_code=0,
                ),
            )

            judge_workdir = root / "judge" / "workdir"
            context = judge_gsb_container.prepare(judge_workdir, judge_instance)
            task = judge_gsb_container.build_task(
                judge_instance, workdir=str(judge_workdir)
            )

            self.assertEqual(
                (Path(context["gold_dir"]) / "gold" / "answer.txt").read_text(
                    encoding="utf-8"
                ),
                "gold",
            )
            self.assertIn("reverse: A is GOLD_DIR, B is CANDIDATE_DIR.", task)
            self.assertIn("forward: A is CANDIDATE_DIR, B is GOLD_DIR.", task)
            self.assertIn(
                str(judge_workdir / judge_gsb_container.JUDGE_RESULT_FILE), task
            )

    def test_judge_suite_local_process_oracle_provider(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate_workdir = root / "candidate_workdir"
            candidate_workdir.mkdir()
            (candidate_workdir / "answer.txt").write_text("answer", encoding="utf-8")
            candidate_result = container.extract_result(
                candidate_workdir, {"task_id": "task-1"}
            )
            suite = GdpvalJudgeSuite(image="unused", network_mode=None)
            judge_instance = suite.build_instance(
                {
                    "task_id": "task-1",
                    "prompt": "Make answer.",
                    "rubrics": [{"score": 2, "criterion": "Answer exists"}],
                },
                candidate_result=candidate_result,
                candidate_artifacts=RunArtifacts(
                    instance_id="task-1",
                    run_dir=root / "solver-run",
                    trajectory_path=root / "solver-run" / "out" / "trajectory.jsonl",
                    status_code=0,
                ),
            )

            artifacts = run_suite_instance(
                suite=suite,
                instance=judge_instance,
                backend=LocalProcessBackend(workspace=root / "judge_workdir"),
                store=LocalDirStore(root / "runs"),
                run_root=root / "runs",
                run_id="judge-smoke",
                provider="oracle",
                max_turns=1,
            )

            self.assertEqual(artifacts.status_code, 0)
            result = json.loads((artifacts.run_dir / RESULT_KEY).read_text())
            self.assertEqual(result["status"], "judged")
            self.assertEqual(result["score"], 1.0)
            self.assertEqual(result["earned_score"], 2.0)

    def test_gsb_judge_suite_local_process_oracle_provider(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate_workdir = root / "candidate_workdir"
            candidate_workdir.mkdir()
            (candidate_workdir / "answer.txt").write_text("answer", encoding="utf-8")
            candidate_result = container.extract_result(
                candidate_workdir, {"task_id": "task-1"}
            )
            deliverable_root = root / "deliverables"
            (deliverable_root / "gold").mkdir(parents=True)
            (deliverable_root / "gold" / "answer.txt").write_text(
                "gold", encoding="utf-8"
            )
            suite = GdpvalGsbJudgeSuite(
                image="unused",
                deliverable_root=deliverable_root,
                network_mode=None,
            )
            judge_instance = suite.build_instance(
                {
                    "task_id": "task-1",
                    "prompt": "Make answer.",
                    "rubrics": [{"score": 2, "criterion": "Answer exists"}],
                    "deliverable_files": ["gold/answer.txt"],
                },
                candidate_result=candidate_result,
                candidate_artifacts=RunArtifacts(
                    instance_id="task-1",
                    run_dir=root / "solver-run",
                    trajectory_path=root / "solver-run" / "out" / "trajectory.jsonl",
                    status_code=0,
                ),
            )

            artifacts = run_suite_instance(
                suite=suite,
                instance=judge_instance,
                backend=LocalProcessBackend(workspace=root / "judge_workdir"),
                store=LocalDirStore(root / "runs"),
                run_root=root / "runs",
                run_id="gsb-judge-smoke",
                provider="oracle",
                max_turns=1,
            )

            self.assertEqual(artifacts.status_code, 0)
            result = json.loads((artifacts.run_dir / RESULT_KEY).read_text())
            self.assertEqual(result["status"], "gsb_judged")
            self.assertEqual(result["judge_mode"], "gsb")
            self.assertEqual(result["score"], 0.5)
            self.assertEqual(result["combined_weighted_score"], 0.5)

    def test_gsb_extract_result_marks_missing_candidate_deliverables(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workdir = root / "judge_workdir"
            workdir.mkdir()
            gold = root / "gold.txt"
            gold.write_text("gold", encoding="utf-8")

            result = judge_gsb_container.extract_result(
                workdir,
                {
                    "task_id": "task-1",
                    "rubrics": [{"score": 2, "criterion": "Answer exists"}],
                },
                context={
                    "candidate_manifest": [],
                    "gold_manifest": [
                        {
                            "name": "gold.txt",
                            "path": str(gold),
                            "missing": False,
                        }
                    ],
                },
            )

            self.assertEqual(result["status"], "candidate_deliverables_missing")
            self.assertEqual(result["score"], 0.0)
            self.assertEqual(result["max_score"], 2.0)
            self.assertIn(
                "no readable candidate deliverable",
                result["overall_explanation_reverse"],
            )

    def test_gsb_extract_result_rejects_missing_forward_direction(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workdir = root / "judge_workdir"
            workdir.mkdir()
            candidate = root / "candidate.txt"
            gold = root / "gold.txt"
            candidate.write_text("candidate", encoding="utf-8")
            gold.write_text("gold", encoding="utf-8")
            (workdir / judge_gsb_container.JUDGE_RESULT_FILE).write_text(
                json.dumps(
                    {
                        "reverse": {
                            "rubrics_result": [
                                {
                                    "score": 1,
                                    "criterion": "Answer exists",
                                    "grade_A": 1,
                                    "grade_B": 0,
                                    "gsb": "A>B",
                                }
                            ],
                            "overall": {
                                "overall_explanation": "gold is better",
                                "final_gsb": "A>B",
                            },
                        }
                    }
                )
                + "\n",
                encoding="utf-8",
            )

            result = judge_gsb_container.extract_result(
                workdir,
                {
                    "task_id": "task-1",
                    "rubrics": [{"score": 1, "criterion": "Answer exists"}],
                },
                context={
                    "candidate_manifest": [
                        {
                            "name": "candidate.txt",
                            "path": str(candidate),
                            "missing": False,
                        }
                    ],
                    "gold_manifest": [
                        {
                            "name": "gold.txt",
                            "path": str(gold),
                            "missing": False,
                        }
                    ],
                },
            )

            self.assertEqual(result["status"], "judge_result_invalid")
            self.assertEqual(result["score"], 0.0)
            self.assertIn(
                "forward judge payload is missing",
                result["overall_explanation_reverse"],
            )

    def test_gsb_judge_fake_provider_records_no_tool_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate_workdir = root / "candidate_workdir"
            candidate_workdir.mkdir()
            (candidate_workdir / "answer.txt").write_text("answer", encoding="utf-8")
            candidate_result = container.extract_result(
                candidate_workdir, {"task_id": "task-1"}
            )
            deliverable_root = root / "deliverables"
            (deliverable_root / "gold").mkdir(parents=True)
            (deliverable_root / "gold" / "answer.txt").write_text(
                "gold", encoding="utf-8"
            )
            suite = GdpvalGsbJudgeSuite(
                image="unused",
                deliverable_root=deliverable_root,
                network_mode=None,
                judge_tool_mode="local",
            )
            judge_instance = suite.build_instance(
                {
                    "task_id": "task-1",
                    "prompt": "Make answer.",
                    "rubrics": [{"score": 1, "criterion": "Answer quality"}],
                    "deliverable_files": ["gold/answer.txt"],
                },
                candidate_result=candidate_result,
                candidate_artifacts=RunArtifacts(
                    instance_id="task-1",
                    run_dir=root / "solver-run",
                    trajectory_path=root / "solver-run" / "out" / "trajectory.jsonl",
                    status_code=0,
                ),
            )

            artifacts = run_suite_instance(
                suite=suite,
                instance=judge_instance,
                backend=LocalProcessBackend(workspace=root / "judge_workdir"),
                store=LocalDirStore(root / "runs"),
                run_root=root / "runs",
                run_id="gsb-judge-fake",
                provider="fake",
                max_turns=1,
            )

            self.assertEqual(artifacts.status_code, 0)
            result = json.loads((artifacts.run_dir / RESULT_KEY).read_text())
            self.assertEqual(result["status"], "judge_result_invalid")
            self.assertEqual(result["score"], 0.0)
            retry_summary = result["judge_retry_summary"]
            self.assertEqual(retry_summary["reverse"]["attempts"], 1)
            self.assertNotIn("forward", retry_summary)
            self.assertFalse(retry_summary["reverse"]["had_tool_messages"])
            self.assertIn("reverse judge failed", result["overall_explanation_reverse"])
            self.assertEqual(
                result["judge_attempts"][0]["failure_reason"], "no_tool_messages"
            )

    def test_run_suite_instance_local_process_fake_provider(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "runs"
            workdir = Path(tmp) / "workspace"
            suite = GdpvalSuite(
                image="unused",
                reference_root=Path(tmp),
                network_mode=None,
            )
            artifacts = run_suite_instance(
                suite=suite,
                instance={
                    "instance_id": "task-1",
                    "task_id": "task-1",
                    "prompt": "Create any file under WORKDIR.",
                },
                backend=LocalProcessBackend(workspace=workdir),
                store=LocalDirStore(root),
                run_root=root,
                run_id="smoke",
                provider="fake",
                max_turns=2,
            )

            self.assertEqual(artifacts.status_code, 0)
            result = json.loads((artifacts.run_dir / RESULT_KEY).read_text())
            self.assertEqual(result["status"], "solver_finished")
            self.assertFalse((artifacts.run_dir / "input" / "eval.json").exists())
            trace = json.loads((artifacts.run_dir / TRACE_KEY).read_text())
            self.assertEqual(trace["meta"]["suite"], "gdpval")


if __name__ == "__main__":
    unittest.main()
